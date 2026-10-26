"""Reads and writes a Medvind base schedule (grundschema) exported to Excel.

The layout, as Medvind exports it: row 1 has the headers "Signaler", "Schemarad", "Attesterad/godkänd t.o.m"
and one merged "Vecka N" header per rotation week (the start week is marked "Vecka 4: Startvecka"); row 2 has the
weekday names, Måndag to Söndag, under each week. From row 3, one row per schedule line (schemarad), e.g.
"1, Avd 1" or "13 Extra", with one time range per cell. The times are typed by hand, so the formats vary:
"06.00-13.00", "8-16.30", "0630-1030", "9-19", and typos such as "6.3014.0".

    read(path)            -> rotation (a dict, see below), with every cell it could not read flagged
    write(rotation, path) -> the same layout, times written as "06.00-13.00"

    python3 medvind/medvind.py read  grundschema.xlsx > rotation.json
    python3 medvind/medvind.py write rotation.json ut.xlsx

rotation = {"weeks": N, "start": k (the start week, 1..N), "lines": [{"row", "label", "no", "unit", "extra",
            "cells": [week][day 0-6] -> {"from", "to"} in minutes, plus "raw", and "flag"/"note" if guessed} or None]}],
            "flags": [{"cell", "row", "line", "week", "day", "raw", "kind", "note"}]}
kind is "guessed" (read, but check it), "unread" (left out) or "noUnit" (a line without a unit).
"""

import json
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DAYS = ["Måndag", "Tisdag", "Onsdag", "Torsdag", "Fredag", "Lördag", "Söndag"]
EARLIEST, LATEST = 5 * 60, 21 * 60  # a time outside 05:00-21:00 is not believed


def _time(text):
    """'6' '06' '6.30' '06:30' '6,30' '0630' -> minutes, or None."""
    m = re.fullmatch(r"(\d{1,2})(?:[.:,](\d{2}))?", text) or re.fullmatch(r"(\d{2})(\d{2})", text)
    if not m:
        return None
    h, mm = int(m.group(1)), int(m.group(2) or 0)
    return h * 60 + mm if h < 24 and mm < 60 else None


def parse_range(raw):
    """One cell -> {"from", "to", "raw"} (with "flag": "guessed" and a "note" if a typo was read), or
    {"raw", "flag": "unread", "note"}, or None for an empty cell."""
    text = re.sub(r"\s+", "", str(raw or "")).replace("–", "-").replace("—", "-")
    if not text:
        return None
    out = {"raw": str(raw).strip()}
    parts = text.split("-")
    if len(parts) == 2:
        a, b = _time(parts[0]), _time(parts[1])
    else:
        a = b = None
        # A missing dash, as in "6.3014.0": a time with minutes followed by another time.
        m = re.fullmatch(r"(\d{1,2}[.:,]\d{2})(\d{1,2})(?:[.:,](\d{1,2}))?", text)
        if m:
            a, b = _time(m.group(1)), int(m.group(2)) * 60 + (int(m.group(3)) * (10 if len(m.group(3)) == 1 else 1) if m.group(3) else 0)
            if a is not None and b < 24 * 60:
                out.update({"flag": "guessed", "note": f"Bindestreck saknas; läst som {_fmt(a)}–{_fmt(b)}."})
            else:
                a = b = None
    if a is None or b is None:
        return {**out, "flag": "unread", "note": "Tiden gick inte att läsa."}
    if not (EARLIEST <= a < b <= LATEST):
        return {**out, "flag": "unread", "note": f"Orimlig tid: {_fmt(a)}–{_fmt(b)}."}
    return {**out, "from": a, "to": b}


def _fmt(m):
    return f"{m // 60:02d}:{m % 60:02d}"


def _medvind(m):
    return f"{m // 60:02d}.{m % 60:02d}"


def read(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    head = {c.column: str(c.value).strip() for c in ws[1] if c.value not in (None, "")}
    line_col = next((c for c, v in head.items() if v.lower().startswith("schemarad")), None)
    weeks = sorted((int(re.search(r"\d+", v).group()), c, "startvecka" in v.lower())
                   for c, v in head.items() if re.match(r"vecka\s*\d+", v, re.I))
    if not line_col or not weeks:
        raise ValueError("Hittar inte kolumnerna Schemarad och Vecka 1, 2 … på första raden. Är det en export av ett grundschema från Medvind?")
    start = next((n for n, _, s in weeks if s), 1)
    lines, flags = [], []
    for r in range(3, ws.max_row + 1):
        label = str(ws.cell(r, line_col).value or "").strip()
        if not label:
            continue
        m = re.match(r"(\d+)", label)
        # The unit is what follows the comma ("1, Avd 1" -> "Avd 1"); a bare "Avd" has no number, so no unit.
        unit = label.split(",", 1)[1].strip() if "," in label else ""
        unit = None if not unit or re.fullmatch(r"avd\.?", unit, re.I) else unit
        extra = "extra" in label.lower()
        line = {"row": r, "label": label, "no": int(m.group(1)) if m else None, "unit": unit, "extra": extra, "cells": []}
        if not line["unit"] and not extra:
            flags.append({"cell": f"{get_column_letter(line_col)}{r}", "row": r, "line": label, "week": None, "day": None, "raw": label,
                          "kind": "noUnit", "note": "Schemaraden har ingen avdelning."})
        for n, col, _ in weeks:
            week = []
            for d in range(7):
                cell = ws.cell(r, col + d)
                x = parse_range(cell.value)
                if x and x.get("flag"):
                    flags.append({"cell": cell.coordinate, "row": r, "line": label, "week": n, "day": d, "raw": x["raw"], "kind": x["flag"], "note": x["note"]})
                week.append(x if x and "from" in x else None)
            line["cells"].append(week)
        lines.append(line)
    return {"weeks": len(weeks), "start": start, "lines": lines, "flags": flags}


def write(rot, path):
    """Writes the rotation in Medvind's layout: the same headers, merged week headers, the start week marked."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medvind"
    bold, center = Font(bold=True), Alignment(horizontal="center", vertical="center")
    for col, name in ((3, "Signaler"), (4, "Schemarad"), (5, "Attesterad/godkänd t.o.m")):
        ws.cell(1, col, name).font = bold
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for c in (1, 2):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for n in range(1, rot["weeks"] + 1):
        col = 6 + (n - 1) * 7
        cell = ws.cell(1, col, f"Vecka {n}" + (": Startvecka" if n == rot.get("start", 1) else ""))
        cell.font, cell.alignment = bold, center
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 6)
        for d, name in enumerate(DAYS):
            ws.cell(2, col + d, name).alignment = center
    flagged = PatternFill("solid", fgColor="FFE3B8")
    for i, line in enumerate(rot["lines"]):
        r = 3 + i
        ws.cell(r, 3, "")
        ws.cell(r, 4, line["label"])
        ws.cell(r, 5, "")
        for n, week in enumerate(line["cells"]):
            for d, x in enumerate(week):
                cell = ws.cell(r, 6 + n * 7 + d, f"{_medvind(x['from'])}-{_medvind(x['to'])}" if x else "")
                if x and x.get("flag"):
                    cell.fill = flagged
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    for c in range(6, 6 + 7 * rot["weeks"]):
        ws.column_dimensions[get_column_letter(c)].width = 12
    wb.save(path)


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "read":
        json.dump(read(sys.argv[2]), sys.stdout, ensure_ascii=False, indent=1)
    elif len(sys.argv) == 4 and sys.argv[1] == "write":
        with open(sys.argv[2], encoding="utf-8") as f:
            write(json.load(f), sys.argv[3])
    else:
        sys.exit(__doc__)
